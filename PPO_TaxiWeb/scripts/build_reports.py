"""Regenerate Word reports from the Markdown specifications (requires python-docx, Pillow)."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt,Cm,Inches
from docx.oxml import OxmlElement
from PIL import Image
root=Path(__file__).resolve().parents[1]
out=root/'docs'/'reports';out.mkdir(exist_ok=True)
def clean(s):return s.replace('**','').replace('`','')
def markdown(doc,path):
 lines=path.read_text().splitlines();i=0;code=False
 while i<len(lines):
  line=lines[i];i+=1
  if line.startswith('```'):code=not code;continue
  if code:
   p=doc.add_paragraph();r=p.add_run(line);r.font.name='Courier New';r.font.size=Pt(9);continue
  if not line.strip():continue
  if line.startswith('|'):
   rows=[line]
   while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
   values=[[clean(x.strip()) for x in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[\s:|\-]+\|$',r)]
   cols=max(map(len,values));t=doc.add_table(rows=0,cols=cols);t.style='Table Grid'
   for ri,row in enumerate(values):
    cells=t.add_row().cells
    for c,v in zip(cells,row):
     c.text=v
     for paragraph in c.paragraphs:
      for run in paragraph.runs:run.font.size=Pt(10);run.bold=(ri==0)
    if ri==0:
     prop=t.rows[0]._tr.get_or_add_trPr();prop.append(OxmlElement('w:tblHeader'))
   doc.add_paragraph();continue
  if line.startswith('#'):
   match=re.match(r'^(#+)\s+(.*)',line)
   if match:doc.add_heading(clean(match[2]),level=min(len(match[1])+0,3));continue
  if line.startswith('- '):doc.add_paragraph(clean(line[2:]),style='List Bullet')
  else:doc.add_paragraph(clean(line))
def picture(doc,path,caption):
 if not path.exists():return
 doc.add_page_break();doc.add_heading(caption,level=2)
 with Image.open(path) as im:w,h=im.size
 width=min(6.25,8.0*w/h)
 doc.add_picture(str(path),width=Inches(width));doc.add_paragraph(caption,style='Caption')
groups=[
 (0,'Расчёт стоимости поездки',['docs/project.md','docs/lab0.md','docs/adaptation.md'],[('domain.png','Диаграмма классов')]),
 (1,'Анализ требований',['docs/requirements.md'],[('use-cases.png','Варианты использования')]),
 (2,'Проектирование системы',['docs/design.md','docs/data.md'],[('components.png','Диаграмма компонентов'),('domain.png','Диаграмма классов'),('er.png','ER-модель данных'),('workflow.png','Жизненный цикл заказа')]),
 (3,'Бизнес-логика и тестирование',['docs/lab3.md','docs/requirements.md','docs/validation.md'],[('domain.png','Актуальная диаграмма классов')]),
 (4,'Хранение данных',['docs/lab4.md','docs/data.md','docs/validation.md'],[('er.png','Физическая схема SQLite')]),
 (5,'Веб-интерфейс',['docs/lab5.md','docs/ui.md','README.md','docs/validation.md'],[])]
for number,title,files,diagrams in groups:
 doc=Document();section=doc.sections[0];section.left_margin=Cm(2.2);section.right_margin=Cm(1.5);section.top_margin=section.bottom_margin=Cm(2)
 normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(6)
 doc.add_heading(f'Лабораторная работа №{number}',0);doc.add_paragraph('Проектирование программного обеспечения',style='Subtitle');doc.add_heading('Сайт «Служба такси»',1);doc.add_paragraph(title)
 doc.add_paragraph('Студент: ____________________\nГруппа: ____________________\n2026 г.')
 doc.add_paragraph('Адаптация PPO-Findirector-Labs.docx. Подробное сопоставление требований — docs/adaptation.md. Исходный архив лектора не предоставлен; разработан самостоятельный проект.')
 doc.add_page_break()
 for fi,f in enumerate(files):
  if fi:doc.add_page_break()
  markdown(doc,root/f)
 for name,caption in diagrams:picture(doc,root/'docs'/'diagrams'/name,caption)
 if number==5:
  picture(doc,root/'docs'/'diagrams'/'ui-mockup.png','Проектный макет интерфейса')
  for name,cap in [('client-comparison.png','Сайт: сравнение тарифов клиентом'),('dispatcher-orders.png','Сайт: рабочее место диспетчера'),('driver-trip.png','Сайт: поездка в пути'),('after-restart.png','Сайт: сохранённые заказы после перезапуска')]:picture(doc,root/'docs'/'test-results'/name,cap)
 footer=section.footer.paragraphs[0];footer.add_run('АИС «Служба такси» • ППО • 2026')
 target=out/f'ЛР{number}_Служба_такси.docx';doc.save(target);print(target.name)
